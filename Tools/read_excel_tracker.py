"""
SPINRISE Excel Tracker Reader
Reads all .xlsx files from the extracted attachments folder,
dumps relevant rows (defect IDs, test cases, status columns) to .txt files
"""

import openpyxl
from pathlib import Path

XLSX_DIR = Path(r"D:\SpinriseEmailExtract\attachments\xlsx")

DEFECT_KEYWORDS = [
    "def-fa", "il-", "pr-06", "pr-12", "pr-20", "fc-ex", "fc-br",
    "fa-add", "sp-i", "br-undo", "nv-", "sv-", "pass", "fail",
    "blocked", "open", "closed", "fixed", "pending", "prstatus",
    "ksp_pr", "appuserlabel", "sms", "undo", "foreclosure",
    "cancellation", "first level", "final level", "pr form",
]

def cell_text(cell):
    if cell.value is None:
        return ""
    return str(cell.value).strip()

def row_is_relevant(row_vals):
    combined = " ".join(row_vals).lower()
    return any(k in combined for k in DEFECT_KEYWORDS)

def extract_xlsx(xlsx_path: Path):
    out_path = xlsx_path.with_suffix(".txt")
    lines = []
    try:
        wb = openpyxl.load_workbook(xlsx_path, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            sheet_lines = []
            for row in ws.iter_rows():
                vals = [cell_text(c) for c in row]
                if not any(vals):
                    continue
                if row_is_relevant(vals):
                    sheet_lines.append(" | ".join(v for v in vals if v))
            if sheet_lines:
                lines.append(f"\n{'='*60}")
                lines.append(f"SHEET: {sheet_name}")
                lines.append(f"{'='*60}")
                lines.extend(sheet_lines)
        wb.close()
    except Exception as e:
        lines.append(f"[ERROR reading {xlsx_path.name}: {e}]")

    if lines:
        out_path.write_text("\n".join(lines), encoding="utf-8")
        print(f"  Extracted: {out_path.name} ({len(lines)} lines)")
    else:
        out_path.write_text(f"No relevant rows found in {xlsx_path.name}", encoding="utf-8")
        print(f"  No relevant rows: {xlsx_path.name}")

def main():
    xlsx_files = list(XLSX_DIR.glob("*.xlsx")) + list(XLSX_DIR.glob("*.xls"))
    if not xlsx_files:
        print(f"No Excel files found in {XLSX_DIR}")
        return
    print(f"Found {len(xlsx_files)} Excel file(s):")
    for f in xlsx_files:
        print(f"  Processing: {f.name}")
        extract_xlsx(f)
    print("\nDone.")

if __name__ == "__main__":
    main()
